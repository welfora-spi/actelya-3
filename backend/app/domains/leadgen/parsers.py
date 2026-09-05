"""Lead Generation — parser deterministici per CSV/TSV/XLSX/PDF testuale/
DOCX/TXT. Nessuna esecuzione di macro/script/contenuto incorporato (openpyxl
in modalita' read_only + data_only, mai keep_vba; pypdf estrae solo testo,
mai JavaScript embedded; python-docx legge solo paragrafi/tabelle, mai
campi/macro). Ogni valore estratto e' trattato come DATO, mai come
istruzione: nessuna interpretazione di formule (vedi export.py per la
neutralizzazione in scrittura). Righe/colonne/testo sempre limitati per
evitare un consumo di memoria/tempo non controllato."""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field

from .models import MAX_COLS_PARSED, MAX_ROWS_PARSED, MAX_TEXT_CHARS_EXTRACTED


class ParseError(Exception):
    """File corrotto, protetto da password o non elaborabile in modo sicuro."""


@dataclass
class ParseResult:
    columns: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    testo_libero: str = ""
    avvisi: list[str] = field(default_factory=list)


def _decode_text(content: bytes) -> str:
    try:
        from charset_normalizer import from_bytes
        migliore = from_bytes(content).best()
        if migliore is not None:
            return str(migliore)
    except Exception:
        pass
    return content.decode("utf-8", errors="replace")


def _parse_delimited(testo: str, delimiter: str | None) -> ParseResult:
    avvisi: list[str] = []
    if delimiter is None:
        try:
            dialect = csv.Sniffer().sniff(testo[:5000], delimiters=",;\t|")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = ","
            avvisi.append("Delimitatore non rilevato automaticamente: uso ',' di default.")
    reader = csv.reader(io.StringIO(testo), delimiter=delimiter)
    grezzo = list(reader)
    grezzo = [r for r in grezzo if any((c or "").strip() for c in r)]  # righe interamente vuote
    if not grezzo:
        return ParseResult([], [], testo, avvisi + ["File vuoto o senza righe valide."])
    header = [c.strip() for c in grezzo[0][:MAX_COLS_PARSED]]
    corpo = grezzo[1:]
    if len(corpo) > MAX_ROWS_PARSED:
        avvisi.append(f"Righe troncate a {MAX_ROWS_PARSED} (il file ne contiene di piu').")
        corpo = corpo[:MAX_ROWS_PARSED]
    corpo = [[c.strip() for c in r[:MAX_COLS_PARSED]] for r in corpo]
    return ParseResult(header, corpo, "", avvisi)


def parse_csv(content: bytes) -> ParseResult:
    return _parse_delimited(_decode_text(content), delimiter=None)


def parse_tsv(content: bytes) -> ParseResult:
    return _parse_delimited(_decode_text(content), delimiter="\t")


def parse_xlsx(content: bytes) -> ParseResult:
    try:
        import openpyxl
    except ImportError as exc:
        raise ParseError("Supporto XLSX non disponibile su questo server.") from exc
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_vba=False)
    except Exception as exc:
        raise ParseError(f"File XLSX non leggibile, protetto o corrotto: {exc}") from exc
    avvisi: list[str] = []
    try:
        ws = wb.worksheets[0]
        righe: list[list[str]] = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > MAX_ROWS_PARSED:
                avvisi.append(f"Righe troncate a {MAX_ROWS_PARSED} (il foglio ne contiene di piu').")
                break
            righe.append([("" if v is None else str(v)).strip() for v in row[:MAX_COLS_PARSED]])
    finally:
        wb.close()
    righe = [r for r in righe if any(c for c in r)]
    if not righe:
        return ParseResult([], [], "", avvisi + ["Il primo foglio del file e' vuoto."])
    return ParseResult(righe[0], righe[1:], "", avvisi)


def _tenta_struttura_da_testo(testo: str) -> tuple[list[str], list[list[str]]]:
    """Euristica leggera e conservativa: SOLO se il testo assomiglia a righe
    con lo stesso delimitatore ripetuto in modo consistente viene trattato
    come tabellare. Mai una struttura forzata su testo libero genuino
    (nessun crawler/parser complesso: se non e' chiaramente tabellare, resta
    solo testo)."""
    righe_testo = [r for r in testo.splitlines() if r.strip()][:MAX_ROWS_PARSED + 1]
    if len(righe_testo) < 2:
        return [], []
    for delim in ("\t", ";", ","):
        conteggi = [r.count(delim) for r in righe_testo[:20]]
        if conteggi and conteggi[0] >= 1 and len(set(conteggi)) == 1:
            grezzo = list(csv.reader(righe_testo, delimiter=delim))
            header = [c.strip() for c in grezzo[0][:MAX_COLS_PARSED]]
            corpo = [[c.strip() for c in r[:MAX_COLS_PARSED]] for r in grezzo[1:MAX_ROWS_PARSED + 1]]
            return header, corpo
    return [], []


def parse_pdf(content: bytes) -> ParseResult:
    try:
        import pypdf
    except ImportError as exc:
        raise ParseError("Supporto PDF non disponibile su questo server.") from exc
    try:
        reader = pypdf.PdfReader(io.BytesIO(content))
    except Exception as exc:
        raise ParseError(f"PDF non leggibile o corrotto: {exc}") from exc
    if reader.is_encrypted:
        raise ParseError("PDF protetto da password: non supportato.")
    avvisi: list[str] = []
    parti: list[str] = []
    for pagina in reader.pages:
        try:
            parti.append(pagina.extract_text() or "")
        except Exception:
            avvisi.append("Una pagina del PDF non e' stata estratta correttamente.")
    testo = "\n".join(parti)[:MAX_TEXT_CHARS_EXTRACTED]
    if not testo.strip():
        avvisi.append("Nessun testo estraibile (PDF probabilmente scansionato/immagine): OCR non supportato.")
    colonne, righe = _tenta_struttura_da_testo(testo)
    return ParseResult(colonne, righe, testo, avvisi)


def parse_docx(content: bytes) -> ParseResult:
    try:
        import docx
    except ImportError as exc:
        raise ParseError("Supporto DOCX non disponibile su questo server.") from exc
    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:
        raise ParseError(f"DOCX non leggibile, protetto o corrotto: {exc}") from exc
    testo = "\n".join(p.text for p in document.paragraphs)[:MAX_TEXT_CHARS_EXTRACTED]
    colonne: list[str] = []
    righe: list[list[str]] = []
    if document.tables:
        dati = [[cell.text.strip() for cell in row.cells] for row in document.tables[0].rows]
        dati = [r for r in dati if any(c for c in r)]
        if len(dati) >= 2:
            colonne, righe = dati[0][:MAX_COLS_PARSED], [r[:MAX_COLS_PARSED] for r in dati[1:MAX_ROWS_PARSED + 1]]
    if not colonne:
        colonne, righe = _tenta_struttura_da_testo(testo)
    return ParseResult(colonne, righe, testo, [])


def parse_txt(content: bytes) -> ParseResult:
    testo = _decode_text(content)[:MAX_TEXT_CHARS_EXTRACTED]
    colonne, righe = _tenta_struttura_da_testo(testo)
    return ParseResult(colonne, righe, testo, [])


_PARSERS = {
    ".csv": parse_csv, ".tsv": parse_tsv, ".xlsx": parse_xlsx,
    ".pdf": parse_pdf, ".docx": parse_docx, ".txt": parse_txt,
}


def parse_file(extension: str, content: bytes) -> ParseResult:
    fn = _PARSERS.get(extension)
    if fn is None:
        raise ParseError(f"Formato non supportato: '{extension}'.")
    return fn(content)
